# scripts/05_export_excel.py
"""
Step 5: Excel workbook.

Writes outputs/parcel_locker_analysis.xlsx with six sheets:
  Sites                  chosen sites in rank order, with priority tier
  Coverage Curve         one row per number of sites
  Districts              access at 10 sites and at the final N, per district
  Candidates             every candidate site, with its chosen rank if chosen
  Headline               the figures the page quotes
  Sources & Assumptions  sources, access dates, thresholds and caveats

Formatting follows RoutePlanner: a single accent colour on a bold frozen header row,
autofilter, explicit column widths, correct number formats, at most one colour
scale per sheet, and A4 landscape print setup with a repeating header row. No
openpyxl Table objects, so the file behaves consistently in LibreOffice Calc.
"""

import geopandas as gpd
import numpy as np
import pandas as pd
import pyogrio
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from config import (BAND_EDGES_MIN, GRID_GPKG, MAX_SITES, OUT_XLSX, SELECTION_GPKG,
                    SITE_MERGE_M, SITES_GPKG, SNAP_EXCLUDE_M, SNAP_FLAG_M, SNAPSHOT_SITES,
                    TARGET_MIN, TARGET_SHARE, TIER_MILESTONES, WALK_CAP_MIN,
                    WALK_SPEED_M_PER_MIN)

# One accent colour, used on header rows. ARGB with an explicit opaque alpha so
# the fill renders solid in LibreOffice as well as Excel.
ACCENT = "FF1F4E79"
SCALE_LOW = "FFFFFFFF"
SCALE_HIGH = "FF9DC3E6"
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)

COUNT = "#,##0"
MINUTES = "0.0"
PERCENT = "0.0%"

# Access dates of the source data (all fetched by script on this date).
ACCESSED = "2026-09-25"


def style_data_sheet(ws, n_cols, n_rows, widths, number_formats):
    """Header styling, widths, number formats, freeze, autofilter and print setup.
    number_formats maps 1-based column index to a format string."""
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    for col, fmt in number_formats.items():
        for row in range(2, n_rows + 2):
            ws.cell(row=row, column=col).number_format = fmt

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows + 1}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"


def colour_scale(ws, column_letter, n_rows):
    ws.conditional_formatting.add(
        f"{column_letter}2:{column_letter}{n_rows + 1}",
        ColorScaleRule(start_type="min", start_color=SCALE_LOW,
                       end_type="max", end_color=SCALE_HIGH),
    )


def weighted_access(cells, minutes_col):
    """Residents, average capped walk and share within TARGET_MIN for a group."""
    res = cells["residents"]
    return pd.Series({
        "residents": int(res.sum()),
        "avg": float(np.average(cells[minutes_col], weights=res)),
        "share": float(res[cells[minutes_col] <= TARGET_MIN].sum() / res.sum()),
    })


# --- Sheets ------------------------------------------------------------------

def build_sites(ws, sites):
    headers = [
        "Rank", "Tier", "Site", "Chain", "District",
        "Residents newly served", "Resident-minutes saved",
        "Average walk after (min)", f"Share within {TARGET_MIN} min after",
        f"Residents nearest at {len(sites)} sites",
    ]
    ws.append(headers)
    for _, s in sites.iterrows():
        ws.append([
            int(s["rank"]), int(s["tier"]), s["name"], s["chain"], s["district"],
            int(s["residents_new"]), round(float(s["resident_min_saved"])),
            float(s["avg_walk_min"]), float(s[f"share_{TARGET_MIN}min"]),
            int(s["residents_nearest_final"]),
        ])
    n_rows = len(sites)
    style_data_sheet(
        ws, len(headers), n_rows,
        widths=[7, 6, 28, 16, 18, 14, 15, 13, 14, 15],
        number_formats={6: COUNT, 7: COUNT, 8: MINUTES, 9: PERCENT, 10: COUNT},
    )
    colour_scale(ws, "G", n_rows)
    return n_rows


def build_curve(ws, curve, tier_by_n):
    headers = ["Sites", "Tier of last site", "Average walk (min)",
               "Share within 5 min", "Share within 10 min", "Share within 15 min"]
    ws.append(headers)
    for _, r in curve.iterrows():
        ws.append([
            int(r["n_sites"]), int(tier_by_n[int(r["n_sites"])]),
            float(r["avg_walk_min"]), float(r["share_5min"]),
            float(r["share_10min"]), float(r["share_15min"]),
        ])
    n_rows = len(curve)
    style_data_sheet(
        ws, len(headers), n_rows,
        widths=[8, 10, 12, 12, 12, 12],
        number_formats={3: MINUTES, 4: PERCENT, 5: PERCENT, 6: PERCENT},
    )
    colour_scale(ws, "E", n_rows)
    return n_rows


def build_districts(ws, cells, candidates, sites, n_final):
    headers = [
        "District", "Residents", "Candidate sites", f"Sites chosen (of {n_final})",
        f"Average walk, {SNAPSHOT_SITES} sites (min)",
        f"Share within {TARGET_MIN} min, {SNAPSHOT_SITES} sites",
        f"Average walk, {n_final} sites (min)",
        f"Share within {TARGET_MIN} min, {n_final} sites",
    ]
    ws.append(headers)
    at_10 = cells.groupby("district").apply(weighted_access, "min_10", include_groups=False)
    at_n = cells.groupby("district").apply(weighted_access, "min_final", include_groups=False)
    n_cand = candidates["district"].value_counts()
    n_chosen = sites["district"].value_counts()
    order = at_10.sort_values("residents", ascending=False).index
    for d in order:
        ws.append([
            d, int(at_10.loc[d, "residents"]), int(n_cand.get(d, 0)), int(n_chosen.get(d, 0)),
            at_10.loc[d, "avg"], at_10.loc[d, "share"], at_n.loc[d, "avg"], at_n.loc[d, "share"],
        ])
    n_rows = len(order)
    style_data_sheet(
        ws, len(headers), n_rows,
        widths=[20, 12, 11, 11, 14, 14, 14, 14],
        number_formats={2: COUNT, 3: COUNT, 4: COUNT, 5: MINUTES, 6: PERCENT,
                        7: MINUTES, 8: PERCENT},
    )
    colour_scale(ws, "H", n_rows)

    # Study-area total, two rows below the table so it stays outside the filter.
    t10 = weighted_access(cells, "min_10")
    tn = weighted_access(cells, "min_final")
    row = n_rows + 3
    values = ["Espoo + Kauniainen", int(t10["residents"]), len(candidates), len(sites),
              t10["avg"], t10["share"], tn["avg"], tn["share"]]
    formats = [None, COUNT, COUNT, COUNT, MINUTES, PERCENT, MINUTES, PERCENT]
    for col, (value, fmt) in enumerate(zip(values, formats), start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(bold=True)
        if fmt:
            cell.number_format = fmt
    return n_rows


def build_candidates(ws, candidates, sites):
    headers = ["Site ID", "Site", "Chain", "District", "Chosen rank", "Tier",
               "Stores merged into this site", "OSM reference"]
    ws.append(headers)
    chosen = sites.set_index("site_id")
    cands = candidates.copy()
    cands["rank"] = cands["site_id"].map(chosen["rank"])
    cands = cands.sort_values(["rank", "district", "name"], na_position="last")
    for _, c in cands.iterrows():
        chosen_row = not pd.isna(c["rank"])
        ws.append([
            int(c["site_id"]), c["name"], c["chain"], c["district"],
            int(c["rank"]) if chosen_row else None,
            int(chosen.loc[c["site_id"], "tier"]) if chosen_row else None,
            c["merged_names"] if isinstance(c["merged_names"], str) and c["merged_names"] else None,
            c["osm_ref"],
        ])
    n_rows = len(cands)
    style_data_sheet(
        ws, len(headers), n_rows,
        widths=[8, 28, 16, 18, 10, 6, 60, 18],
        number_formats={5: "0", 6: "0"},
    )
    for row in range(2, n_rows + 2):
        ws.cell(row=row, column=7).alignment = Alignment(wrap_text=True, vertical="top")
    return n_rows


def build_text_sheet(ws, headers, rows, widths):
    """Two- or three-column reference sheet. Each row is a list of values; a value
    given as (value, number_format) is formatted."""
    ws.append(headers)
    for row in rows:
        values = [v[0] if isinstance(v, tuple) else v for v in row]
        ws.append(values)
        r = ws.max_row
        for col, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(v, tuple):
                cell.number_format = v[1]
                cell.alignment = Alignment(horizontal="right", vertical="top")
    style_data_sheet(ws, len(headers), len(rows), widths=widths, number_formats={})
    ws.auto_filter.ref = None
    return len(rows)


def headline_rows(summary, curve, sites, tier_table):
    at = curve.set_index("n_sites")
    n = int(summary["n_final"])
    first = sites.iloc[0]
    chains = sites["chain"].value_counts()
    rows = [
        ["Residents in Espoo + Kauniainen (HSY grid 2025)", (int(summary["residents"]), COUNT), ""],
        ["Populated 250 m cells", (int(summary["cells"]), COUNT), ""],
        ["Candidate sites", (int(summary["candidates"]), COUNT),
         "S Group, K Group, Lidl, R-kioski stores and shopping centres"],
        [f"Best possible share within {TARGET_MIN} min (all candidates open)",
         (summary["ceiling_share"], PERCENT), "The ceiling for these host types"],
        ["Average walk with all candidates open (min)",
         (summary["ceiling_avg_walk_min"], MINUTES), ""],
        [f"Target: {TARGET_SHARE:.0%} of the ceiling", (summary["target_share"], PERCENT), ""],
        ["Sites needed to reach the target", (n, COUNT), ""],
        [f"Share within {TARGET_MIN} min at {SNAPSHOT_SITES} sites",
         (at.loc[SNAPSHOT_SITES, f"share_{TARGET_MIN}min"], PERCENT), ""],
        [f"Average walk at {SNAPSHOT_SITES} sites (min)",
         (at.loc[SNAPSHOT_SITES, "avg_walk_min"], MINUTES), ""],
        [f"Share within {TARGET_MIN} min at {n} sites", (at.loc[n, f"share_{TARGET_MIN}min"], PERCENT), ""],
        [f"Share within 15 min at {n} sites", (at.loc[n, "share_15min"], PERCENT), ""],
        [f"Average walk at {n} sites (min)", (at.loc[n, "avg_walk_min"], MINUTES), ""],
        ["First site", first["name"],
         f"{int(first['residents_new']):,} residents newly served"],
        ["Resident-minutes saved, site 1 / site "
         f"{n}", f"{first['resident_min_saved']:,.0f} / {sites.iloc[-1]['resident_min_saved']:,.0f}",
         "Diminishing returns"],
        ["Chosen sites by chain", "; ".join(f"{c} {k}" for c, k in chains.items()), ""],
    ]
    for t in tier_table:
        rows.append([f"Tier {t['tier']}: sites {t['ranks']}",
                     (t["share"], PERCENT),
                     f"{t['sites']} sites; {t['of_ceiling']:.0%} of the ceiling"])
    return rows


def sources_rows(summary):
    bands = ", ".join([f"under {BAND_EDGES_MIN[0]}"]
                      + [f"{a}-{b}" for a, b in zip(BAND_EDGES_MIN[:-1], BAND_EDGES_MIN[1:])]
                      + [f"over {BAND_EDGES_MIN[-1]}"])
    return [
        ["Source", "Population grid 2025 (Vaestotietoruudukko_2025), municipal and major "
         "district boundaries, rail and metro lines, stations, sea area. Helsinki Region "
         f"Environmental Services HSY, open WFS, CC BY 4.0. Accessed {ACCESSED}."],
        ["Source", "Candidate stores and shopping centres: OpenStreetMap contributors via "
         f"the Overpass API, ODbL. Accessed {ACCESSED}."],
        ["Source", "Walking network: OpenStreetMap contributors via OSMnx, ODbL. "
         f"Accessed {ACCESSED}."],
        ["Source", "Official population for the grid check: Statistics Finland, table "
         f"11ra. Accessed {ACCESSED}."],
        ["Source", "R-kioski store list used to exclude closed kiosks: r-kioski.fi/kioskit. "
         f"Accessed {ACCESSED}."],
        ["Site types", "Candidate sites are limited to the location types Budbee currently "
         "uses to host its lockers: S Group, K Group, Lidl and R-kioski stores, plus "
         "shopping centres. No Budbee data is used; this is a from-scratch design, not a "
         "description of any operator's network."],
        ["Method", "Walking distance along the network from each populated cell centre "
         "to each candidate, plus the straight-line snap to the network at both ends. "
         "Sites added one at a time, each the one that most reduces total resident "
         "walking time."],
        ["Threshold", f"Walking speed {WALK_SPEED_M_PER_MIN} m per minute (4.8 km/h)"],
        ["Threshold", f"Walk cap {WALK_CAP_MIN} minutes: longer walks count as {WALK_CAP_MIN}"],
        ["Threshold", f"Stores within {SITE_MERGE_M} m of each other, or inside or within "
         f"{SITE_MERGE_M} m of a shopping centre outline, are one site"],
        ["Threshold", f"Cells snapped more than {SNAP_FLAG_M} m from the network are "
         f"flagged; more than {SNAP_EXCLUDE_M} m excluded ({int(summary['cells_excluded'])} "
         "cells excluded)"],
        ["Threshold", f"Stop when the share within {TARGET_MIN} minutes reaches "
         f"{TARGET_SHARE:.0%} of the share achievable with every candidate open, or at "
         f"{MAX_SITES} sites"],
        ["Threshold", f"Priority tiers: sites 1-{SNAPSHOT_SITES}, then tiers ending at "
         + ", ".join(f"{m:.0%}" for m in TIER_MILESTONES)
         + f" and {TARGET_SHARE:.0%} of the achievable share"],
        ["Threshold", f"Walking-time bands (minutes): {bands}"],
        ["Assumption", "Network: OSMnx 'walk' filter with shared foot and cycle paths "
         "(highway=cycleway) included, as Finland maps most footpaths that way"],
        ["Caveat", "Residents stand in for demand; pickups near work or on commutes are "
         "not captured."],
        ["Caveat", "Candidate stores come from OpenStreetMap and may be incomplete or "
         "out of date. Three closed R-kioskis and one mis-tagged car dealership were "
         "removed by hand."],
        ["Caveat", "One locker per site, with no capacity limit."],
        ["Caveat", "Lockers are placed only inside Espoo and Kauniainen; lockers across the "
         "Helsinki and Vantaa boundaries are ignored, so edge areas look slightly worse "
         "served than they are."],
        ["Caveat", "250 m grid (about a 3-minute walk across); each cell is measured from "
         "its centre. HSY suppresses age detail in small cells; only total residents are "
         "used."],
    ]


# --- Main --------------------------------------------------------------------

def main():
    cells = gpd.read_file(SELECTION_GPKG, layer="cells_access")
    sites = gpd.read_file(SELECTION_GPKG, layer="sites_selected").sort_values("rank")
    curve = pyogrio.read_dataframe(SELECTION_GPKG, layer="coverage_curve")
    summary = pyogrio.read_dataframe(SELECTION_GPKG, layer="run_summary")
    summary = summary.set_index("item")["value"]
    candidates = gpd.read_file(SITES_GPKG, layer="candidates")
    n_final = int(summary["n_final"])

    tier_by_n = sites.set_index("rank")["tier"].to_dict()
    tier_table = []
    for tier, group in sites.groupby("tier"):
        end = int(group["rank"].max())
        share = float(curve.loc[curve["n_sites"] == end, f"share_{TARGET_MIN}min"].iloc[0])
        tier_table.append({"tier": int(tier),
                           "ranks": f"{int(group['rank'].min())}-{end}",
                           "sites": len(group), "share": share,
                           "of_ceiling": share / summary["ceiling_share"]})

    wb = Workbook()
    ws_sites = wb.active
    ws_sites.title = "Sites"
    ws_curve = wb.create_sheet("Coverage Curve")
    ws_districts = wb.create_sheet("Districts")
    ws_candidates = wb.create_sheet("Candidates")
    ws_headline = wb.create_sheet("Headline")
    ws_sources = wb.create_sheet("Sources & Assumptions")

    n_sites = build_sites(ws_sites, sites)
    n_curve = build_curve(ws_curve, curve, tier_by_n)
    n_districts = build_districts(ws_districts, cells, candidates, sites, n_final)
    n_candidates = build_candidates(ws_candidates, candidates, sites)
    n_headline = build_text_sheet(
        ws_headline, ["Figure", "Value", "Note"],
        headline_rows(summary, curve, sites, tier_table), widths=[52, 26, 50])
    n_sources = build_text_sheet(
        ws_sources, ["Type", "Detail"], sources_rows(summary), widths=[14, 110])

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)

    print("\n--- Summary ---")
    print(f"Wrote {OUT_XLSX}")
    print(f"  Sites                  {n_sites} chosen sites")
    print(f"  Coverage Curve         {n_curve} rows")
    print(f"  Districts              {n_districts} districts plus a total row")
    print(f"  Candidates             {n_candidates} candidates, {len(sites)} chosen")
    print(f"  Headline               {n_headline} figures")
    print(f"  Sources & Assumptions  {n_sources} rows")


if __name__ == "__main__":
    main()
